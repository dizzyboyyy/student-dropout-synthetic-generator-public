from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
from collections import defaultdict, deque
from pathlib import Path
from xml.etree import ElementTree

import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
INPUT_ROOT = ROOT / "inputs" / "demo"
SCENARIO = "public_demo"
SEED = 42

if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from generator_v2.audit import run_structural_audit  # noqa: E402
from generator_v2.curriculum import Curriculum  # noqa: E402
from generator_v2.io_inputs import load_inputs  # noqa: E402
from generator_v2.outputs import (  # noqa: E402
    ANALYSIS_READY_WORKBOOK,
    ANALYSIS_READY_SHEETS,
    verify_run_directory,
)
from generator_v2.validation import validate_inputs  # noqa: E402
from scripts.public_deidentification_gate import scan_roots  # noqa: E402


class PublicReleaseTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.bundle = load_inputs(INPUT_ROOT)
        cls.preflight = validate_inputs(cls.bundle)
        cls.temp = tempfile.TemporaryDirectory(prefix="public_demo_tests_")
        cls.temp_root = Path(cls.temp.name)
        cls.run_one = cls.temp_root / "run_one"
        cls.run_two = cls.temp_root / "run_two"
        cls.runner_outputs: list[str] = []
        for target in (cls.run_one, cls.run_two):
            completed = subprocess.run(
                [
                    sys.executable,
                    "-B",
                    str(ROOT / "scripts" / "run_demo_generator.py"),
                    "--seed",
                    str(SEED),
                    "--output",
                    str(target),
                ],
                cwd=ROOT,
                text=True,
                capture_output=True,
                check=False,
            )
            if completed.returncode != 0:
                raise RuntimeError(
                    f"Demo runner failed ({completed.returncode}):\n"
                    f"STDOUT:\n{completed.stdout}\nSTDERR:\n{completed.stderr}"
                )
            cls.runner_outputs.append(completed.stdout)
        cls.verified = verify_run_directory(
            cls.run_one,
            expected_scenario=SCENARIO,
            expected_seed=SEED,
        )

    @classmethod
    def tearDownClass(cls) -> None:
        cls.temp.cleanup()

    def test_public_imports_and_demo_files(self) -> None:
        self.assertEqual(self.bundle.paths.courses.name, "Demo_Courses.xlsx")
        self.assertEqual(self.bundle.paths.trajectories.name, "Demo_Trajectories.xlsx")
        self.assertEqual(set(self.bundle.hashes), {"Demo_Courses.xlsx", "Demo_Trajectories.xlsx"})

    def test_preflight_49_of_49(self) -> None:
        self.assertEqual(len(self.preflight.checks), 49)
        self.assertTrue(self.preflight.passed, self.preflight.failed)

    def test_demo_curriculum_is_acyclic_and_independent(self) -> None:
        courses = self.bundle.courses
        self.assertEqual(courses["dag_id"].tolist(), [f"COURSE_{i:03d}" for i in range(1, 51)])
        self.assertEqual(courses["nivel"].astype(int).nunique(), 10)
        prerequisites = {
            str(row.codigo_asignatura): list(row.prerequisite_list)
            for row in self.bundle.prerequisites.itertuples(index=False)
        }
        edge_count = sum(len(values) for values in prerequisites.values())
        self.assertNotEqual(edge_count, 69)
        indegree = {node: 0 for node in prerequisites}
        successors: dict[str, list[str]] = defaultdict(list)
        for target, sources in prerequisites.items():
            for source in sources:
                successors[source].append(target)
                indegree[target] += 1
        queue = deque(node for node, degree in indegree.items() if degree == 0)
        visited = 0
        while queue:
            source = queue.popleft()
            visited += 1
            for target in successors[source]:
                indegree[target] -= 1
                if indegree[target] == 0:
                    queue.append(target)
        self.assertEqual(visited, 50)

    def test_demo_admission_is_fictitious_and_complete(self) -> None:
        profiles = self.bundle.profiles
        self.assertEqual(len(profiles), 300)
        self.assertEqual(profiles["id_perfil"].iloc[0], "PROFILE_0001")
        self.assertEqual(profiles["id_perfil"].iloc[-1], "PROFILE_0300")
        self.assertEqual(
            set(profiles["via_ingreso"]),
            {"REGULAR", "ADMISSION_A", "ADMISSION_B", "ADMISSION_C"},
        )

    def test_runner_reports_all_gates(self) -> None:
        output = self.runner_outputs[0]
        self.assertIn("PREFLIGHT PASS: 49/49", output)
        self.assertIn("STRUCTURAL AUDIT PASS: 34/34", output)
        self.assertIn("RUN VERIFICATION PASS", output)

    def test_output_contract_is_complete(self) -> None:
        expected = {
            "Ramos.csv",
            "Estudiantes.csv",
            "Matriculas_estudiante_periodo.csv",
            "Inscripciones_estudiantes.csv",
            "Agregados_curso_periodo.csv",
            "Auditoria.csv",
            "metadata.json",
            "table_hashes.json",
            "run_summary.md",
            ANALYSIS_READY_WORKBOOK,
        }
        self.assertEqual({path.name for path in self.run_one.iterdir() if path.is_file()}, expected)

    def test_structural_audit_34_of_34(self) -> None:
        self.assertEqual(int(self.verified["metadata"]["structural_checks"]["n_checks"]), 34)
        curriculum = Curriculum.from_frames(self.bundle.courses, self.bundle.prerequisites)
        self.assertEqual(curriculum.course_order[-1], "COURSE_050")

    def test_analysis_ready_workbook_contract(self) -> None:
        path = self.run_one / ANALYSIS_READY_WORKBOOK
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            self.assertEqual(workbook.sheetnames, list(ANALYSIS_READY_SHEETS) + ["Data_dictionary"])
            self.assertNotIn("Auditoria", workbook.sheetnames)
        finally:
            workbook.close()

    def test_seed_reproducibility_is_byte_exact_for_six_csvs(self) -> None:
        first = json.loads((self.run_one / "table_hashes.json").read_text(encoding="utf-8"))
        second = json.loads((self.run_two / "table_hashes.json").read_text(encoding="utf-8"))
        self.assertEqual(first, second)
        self.assertEqual(len(first), 6)

    def test_public_deidentification_gate_including_run(self) -> None:
        report = scan_roots([ROOT, self.run_one])
        self.assertTrue(report["passed"], report)
        self.assertEqual(report["forbidden_token_matches"], 0)
        self.assertEqual(report["personal_metadata_matches"], 0)
        self.assertEqual(report["local_path_matches"], 0)
        self.assertEqual(report["structural_privacy_matches"], 0)

    def test_input_and_output_workbook_metadata_is_clean(self) -> None:
        workbooks = [
            INPUT_ROOT / "Demo_Courses.xlsx",
            INPUT_ROOT / "Demo_Trajectories.xlsx",
            self.run_one / ANALYSIS_READY_WORKBOOK,
        ]
        for path in workbooks:
            with self.subTest(path=path.name), zipfile.ZipFile(path) as archive:
                names = [name.casefold() for name in archive.namelist()]
                self.assertFalse(any("comments" in name for name in names))
                self.assertFalse(any("externallinks" in name for name in names))
                self.assertNotIn("docprops/custom.xml", names)
                root = ElementTree.fromstring(archive.read("docProps/core.xml"))
                values = {
                    element.tag.rsplit("}", 1)[-1]: (element.text or "").strip()
                    for element in root.iter()
                }
                self.assertEqual(values.get("creator", ""), "")
                self.assertEqual(values.get("lastModifiedBy", ""), "")


if __name__ == "__main__":
    unittest.main()
