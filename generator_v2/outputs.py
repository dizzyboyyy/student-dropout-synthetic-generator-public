"""Strict, reproducible output persistence and run verification."""

from __future__ import annotations

import hashlib
import json
import math
import re
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import InputBundle
from .engine import SimulationResult
from .schema import AUDIT_ONLY_COLUMN_NAMES, LOGICAL_TABLE_NAMES, RUN_KEYS


class StrictJSONError(ValueError):
    """Raised when a value cannot be represented by strict RFC-compatible JSON."""


class RunVerificationError(ValueError):
    """Raised when a persisted run is incomplete or internally inconsistent."""


ANALYSIS_READY_WORKBOOK = "synthetic_student_data.xlsx"
ANALYSIS_READY_SHEETS = (
    "Ramos",
    "Estudiantes",
    "Matriculas_estudiante_periodo",
    "Inscripciones_estudiantes",
    "Agregados_curso_periodo",
)
DATA_DICTIONARY_COLUMNS = (
    "table",
    "column",
    "role",
    "allowed_as_predictor",
    "temporal_role",
    "description",
)
RELEASE_STATUS = {
    "calibration_status": "NOT_APPLICABLE_DEMO",
    "validation_status": "DEMO_TECHNICAL_PASS",
    "plausibility_status": "NOT_INSTITUTIONALLY_VALIDATED",
    "generator_status": "PUBLIC_SYNTHETIC_DEMO",
}


def _apply_release_reporting_status(metadata: dict[str, Any]) -> None:
    """Record the public-demo status without changing simulated data."""

    validation_status = metadata.get("validation_status")
    if not isinstance(validation_status, dict):
        raise RunVerificationError("metadata.validation_status debe ser un objeto")
    validation_status["calibration"] = "NOT_APPLICABLE_DEMO"
    validation_status["plausibility"] = "DEMONSTRATION_ONLY"
    metadata["release_status"] = RELEASE_STATUS.copy()


def _safe_scenario(value: str) -> str:
    safe = re.sub(r"[^A-Za-z0-9_.-]+", "_", value).strip("._")
    if not safe:
        raise ValueError("scenario_id no produce un nombre de directorio válido")
    return safe


def run_directory(output_root: str | Path, scenario_id: str, seed: int) -> Path:
    """Return the isolated directory assigned to one scenario/seed pair."""

    if int(seed) < 0:
        raise ValueError("La seed debe ser un entero no negativo")
    return Path(output_root).resolve() / f"{_safe_scenario(str(scenario_id))}_seed_{int(seed)}"


def normalize_json_value(value: Any, path: str = "$") -> Any:
    """Normalize supported Python/pandas/numpy values for strict JSON output."""

    if value is None:
        return None
    if isinstance(value, (bool, np.bool_)):
        return bool(value)
    if isinstance(value, (int, np.integer)) and not isinstance(value, bool):
        return int(value)
    if isinstance(value, (float, np.floating)):
        number = float(value)
        if math.isinf(number):
            raise StrictJSONError(f"Valor infinito no permitido en {path}: {number}")
        return None if math.isnan(number) else number
    if isinstance(value, str):
        return value
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, dict):
        normalized: dict[str, Any] = {}
        for key, item in value.items():
            key_text = str(key)
            normalized[key_text] = normalize_json_value(item, f"{path}.{key_text}")
        return normalized
    if isinstance(value, (list, tuple, set)):
        return [normalize_json_value(item, f"{path}[{index}]") for index, item in enumerate(value)]
    if isinstance(value, (pd.Timestamp, np.datetime64)):
        if pd.isna(value):
            return None
        return pd.Timestamp(value).isoformat()
    try:
        missing = pd.isna(value)
    except (TypeError, ValueError):
        missing = False
    if isinstance(missing, (bool, np.bool_)) and bool(missing):
        return None
    raise StrictJSONError(f"Tipo no serializable en {path}: {type(value)!r}")


def strict_json_dumps(value: Any) -> str:
    """Serialize JSON while rejecting all non-standard numeric constants."""

    normalized = normalize_json_value(value)
    return json.dumps(
        normalized,
        ensure_ascii=False,
        indent=2,
        sort_keys=True,
        allow_nan=False,
    )


def _reject_json_constant(value: str) -> None:
    raise StrictJSONError(f"Constante JSON no estándar: {value}")


def strict_json_loads(text: str) -> Any:
    """Parse JSON and reject NaN/Infinity extensions accepted by Python by default."""

    try:
        return json.loads(text, parse_constant=_reject_json_constant)
    except json.JSONDecodeError as exc:
        raise StrictJSONError(f"JSON inválido: {exc}") from exc


def _dataframe_bytes(frame: pd.DataFrame) -> bytes:
    return frame.to_csv(
        index=False,
        lineterminator="\n",
        float_format="%.17g",
    ).encode("utf-8")


def _sha256_bytes(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest().upper()


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def _read_canonical_csv(path: Path) -> pd.DataFrame:
    """Read one generated CSV without treating literal NA-like text as missing."""

    return pd.read_csv(
        path,
        keep_default_na=False,
        na_values=[""],
        low_memory=False,
    )


def _excel_value(value: Any) -> Any:
    """Convert pandas/numpy scalars to openpyxl-compatible typed values."""

    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(value, np.generic):
        return value.item()
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    return value


def _number_format_for(column: str, series: pd.Series) -> str | None:
    """Return display-only formats without changing stored numeric values."""

    if column in {"nota_oficial", "nota_previa"}:
        return "0.0"
    if column in {
        "media",
        "DE",
        "promedio_semestre",
        "promedio_acumulado",
        "promedio_acumulado_pre",
        "promedio_ultimo_semestre_pre",
    }:
        return "0.###"
    if column in {
        "tasa_aprobacion",
        "proporcion_aprobacion",
        "proporcion_aprobacion_pre",
        "senal_admision",
    } or column.startswith("z_"):
        return "0.####"
    if column == "puntaje_admision":
        return "0.##"
    if pd.api.types.is_integer_dtype(series.dtype):
        return "0"
    if pd.api.types.is_float_dtype(series.dtype):
        return "0.####"
    return None


def _write_literal_cell(cell: Cell, value: Any) -> None:
    """Write text beginning with '=' as literal data, never as a formula."""

    cell.value = value
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"


def _style_dataset_sheet(sheet: Any, frame: pd.DataFrame) -> None:
    """Apply restrained, analysis-friendly formatting to one flat table."""

    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Arial", size=10, color="1F1F1F")
    header_alignment = Alignment(horizontal="center", vertical="center")
    text_alignment = Alignment(horizontal="left", vertical="center")
    numeric_alignment = Alignment(horizontal="right", vertical="center")

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 24

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    for column_index, column in enumerate(frame.columns, start=1):
        series = frame[column]
        number_format = _number_format_for(str(column), series)
        alignment = numeric_alignment if number_format is not None else text_alignment
        for row_index in range(2, sheet.max_row + 1):
            cell = sheet.cell(row=row_index, column=column_index)
            cell.font = body_font
            cell.alignment = alignment
            if number_format is not None and cell.value is not None:
                cell.number_format = number_format

        sample_lengths = [len(str(column))]
        for value in series.head(500):
            converted = _excel_value(value)
            if converted is not None:
                sample_lengths.append(len(str(converted)))
        sheet.column_dimensions[get_column_letter(column_index)].width = min(
            max(max(sample_lengths) + 2, 10),
            40,
        )


def _data_dictionary_rows(
    frames: dict[str, pd.DataFrame], metadata: dict[str, Any]
) -> list[list[str]]:
    """Build a column-level dictionary from the persisted metadata contract."""

    contract = metadata.get("data_contract", {})
    roles_by_table = contract.get("public_column_roles", {})
    role_descriptions = contract.get("roles", {})
    predictor_permission = {
        "IDENTIFIER": "NO",
        "REFERENCE": "CONTEXT_DEPENDENT",
        "FEATURE_PRE_T": "YES",
        "OUTCOME_POST_T": "NO",
        "AUDIT_ONLY": "NO",
    }
    temporal_role = {
        "IDENTIFIER": "KEY",
        "REFERENCE": "CONTEXT_DEPENDENT",
        "FEATURE_PRE_T": "PRE_T",
        "OUTCOME_POST_T": "POST_T",
        "AUDIT_ONLY": "INTERNAL_ONLY",
    }
    rows: list[list[str]] = []
    for table in ANALYSIS_READY_SHEETS:
        declared_roles = roles_by_table.get(table, {})
        for column in frames[table].columns:
            role = declared_roles.get(str(column))
            if role not in predictor_permission:
                raise RunVerificationError(
                    f"Data_dictionary no puede resolver el rol de {table}.{column}"
                )
            rows.append(
                [
                    table,
                    str(column),
                    str(role),
                    predictor_permission[str(role)],
                    temporal_role[str(role)],
                    str(role_descriptions.get(str(role), "")),
                ]
            )
    return rows


def write_analysis_ready_workbook(run_dir: str | Path) -> Path:
    """Create the derived relational XLSX exclusively from generated CSV files."""

    directory = Path(run_dir).resolve()
    csv_hashes_before = {
        name: _sha256_file(directory / f"{name}.csv")
        for name in sorted(LOGICAL_TABLE_NAMES)
    }
    metadata = strict_json_loads(
        (directory / "metadata.json").read_text(encoding="utf-8")
    )
    if not isinstance(metadata, dict):
        raise RunVerificationError("metadata.json debe ser un objeto para crear el XLSX")

    frames = {
        name: _read_canonical_csv(directory / f"{name}.csv")
        for name in ANALYSIS_READY_SHEETS
    }
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = ""
    workbook.properties.lastModifiedBy = ""
    workbook.properties.title = "Synthetic student data"
    workbook.properties.subject = "ANALYSIS_READY_RELATIONAL_DATASET"
    workbook.properties.description = (
        "Derived from the five public CSV tables. This workbook is not a final model matrix."
    )

    for table in ANALYSIS_READY_SHEETS:
        frame = frames[table]
        sheet = workbook.create_sheet(table)
        sheet.append([str(column) for column in frame.columns])
        for row_index, values in enumerate(
            frame.itertuples(index=False, name=None),
            start=2,
        ):
            for column_index, value in enumerate(values, start=1):
                _write_literal_cell(
                    sheet.cell(row=row_index, column=column_index),
                    _excel_value(value),
                )
        _style_dataset_sheet(sheet, frame)

    dictionary_sheet = workbook.create_sheet("Data_dictionary")
    dictionary_sheet.append(list(DATA_DICTIONARY_COLUMNS))
    dictionary_rows = _data_dictionary_rows(frames, metadata)
    for row in dictionary_rows:
        dictionary_sheet.append(row)
    dictionary_frame = pd.DataFrame(
        dictionary_rows,
        columns=DATA_DICTIONARY_COLUMNS,
    )
    _style_dataset_sheet(dictionary_sheet, dictionary_frame)
    for column, width in {
        "A": 34,
        "B": 38,
        "C": 20,
        "D": 23,
        "E": 24,
        "F": 58,
    }.items():
        dictionary_sheet.column_dimensions[column].width = width

    workbook_path = directory / ANALYSIS_READY_WORKBOOK
    temporary_path = directory / f".{ANALYSIS_READY_WORKBOOK}.tmp.xlsx"
    try:
        workbook.save(temporary_path)
        temporary_path.replace(workbook_path)
    finally:
        workbook.close()
        if temporary_path.exists():
            temporary_path.unlink()

    csv_hashes_after = {
        name: _sha256_file(directory / f"{name}.csv")
        for name in sorted(LOGICAL_TABLE_NAMES)
    }
    if csv_hashes_after != csv_hashes_before:
        workbook_path.unlink(missing_ok=True)
        raise RunVerificationError("El exportador XLSX modificó uno o más CSV científicos")
    return workbook_path


def verify_analysis_ready_workbook(run_dir: str | Path) -> None:
    """Verify the derived workbook's sheet and leakage-safety contracts."""

    directory = Path(run_dir).resolve()
    workbook_path = directory / ANALYSIS_READY_WORKBOOK
    try:
        workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    except Exception as exc:
        raise RunVerificationError(
            f"No se puede abrir {ANALYSIS_READY_WORKBOOK}: {exc}"
        ) from exc
    try:
        expected_sheets = list(ANALYSIS_READY_SHEETS) + ["Data_dictionary"]
        if workbook.sheetnames != expected_sheets:
            raise RunVerificationError(
                f"Hojas XLSX inválidas: esperadas={expected_sheets}, observadas={workbook.sheetnames}"
            )
        if "Auditoria" in workbook.sheetnames:
            raise RunVerificationError("Auditoria no puede estar incluida en el XLSX")
        for table in ANALYSIS_READY_SHEETS:
            sheet = workbook[table]
            header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
            leaked = sorted(set(header) & AUDIT_ONLY_COLUMN_NAMES)
            if leaked:
                raise RunVerificationError(
                    f"{table}: columnas AUDIT_ONLY expuestas en XLSX {leaked}"
                )
    finally:
        workbook.close()


def build_run_summary(result: SimulationResult, bundle: InputBundle, output_dir: Path) -> str:
    """Build the portable report for one public-demo run."""

    students = result.estudiantes
    panel = result.matriculas
    registrations = result.inscripciones
    aggregates = result.agregados
    audit = result.auditoria
    cohort_sizes = students.groupby("cohorte").size()
    final_states = pd.crosstab(students["cohorte"], students["estado_final"])
    durations = panel.groupby("student_id").size()
    load_counts = panel["carga"].value_counts().sort_index()
    retry_counts = registrations["intento"].value_counts().sort_index()
    course_audit = audit.loc[audit["record_type"].eq("course_state")]
    beta_total = int(course_audit["ajuste_beta"].fillna(0).astype(int).sum())
    beta_denominator = len(course_audit)
    beta_rate = beta_total / beta_denominator if beta_denominator else 0.0
    observed = aggregates.loc[aggregates["N"].gt(0)].copy()
    global_mean = float(registrations["nota_oficial"].mean())
    global_sd = float(registrations["nota_oficial"].std(ddof=0))
    global_pass = float(registrations["aprobado"].mean())
    validation_status = result.metadata["validation_status"]
    release_status = result.metadata["release_status"]

    lines = [
        "# Resumen de corrida — demostración sintética pública",
        "",
        f"- Escenario: `{result.metadata['scenario_id']}`",
        f"- Seed: `{result.metadata['seed']}`",
        f"- Run ID: `{result.metadata['run_id']}`",
        f"- Auditoría estructural: **PASS ({result.structural_checks['n_checks']} controles)**",
        f"- JSON estricto: **{validation_status['strict_json']}**",
        f"- Calibración: **{validation_status['calibration']}**",
        f"- Validación: **{release_status['validation_status']}**",
        f"- Plausibilidad: **{validation_status['plausibility']}**",
        f"- Estado del generador: **{release_status['generator_status']}**",
        f"- Estudiantes: {len(students)}",
        f"- Inscripciones estudiante–curso–período: {len(registrations)}",
        f"- Paneles estudiante–período: {len(panel)}",
        f"- Analysis-ready workbook: `{ANALYSIS_READY_WORKBOOK}`",
        "",
        "## Estudiantes por cohorte",
        "",
        "| Cohorte | n |",
        "|---:|---:|",
    ]
    lines.extend(f"| {int(cohort)} | {int(count)} |" for cohort, count in cohort_sizes.items())
    state_columns = sorted(final_states.columns.astype(str))
    lines.extend(
        [
            "",
            "## Estados finales por cohorte",
            "",
            "| Cohorte | " + " | ".join(state_columns) + " |",
            "|---:|" + "---:|" * len(state_columns),
        ]
    )
    for cohort, row in final_states.iterrows():
        lines.append(
            f"| {int(cohort)} | "
            + " | ".join(str(int(row.get(column, 0))) for column in state_columns)
            + " |"
        )
    lines.extend(
        [
            "",
            "## Trayectorias y notas",
            "",
            "- Número de períodos observados hasta evento terminal o censura: "
            f"media {durations.mean():.2f}, mediana {durations.median():.2f}, "
            f"rango {durations.min()}–{durations.max()}.",
            f"- Nota oficial global: media {global_mean:.3f}, DE poblacional {global_sd:.3f}.",
            f"- Tasa global de aprobación: {global_pass:.3%}.",
            f"- Curso–período con estudiantes: {len(observed)} de {len(aggregates)} contextos.",
            "- El detalle público por curso–período está en `Agregados_curso_periodo.csv`.",
            "",
            "### Distribución de carga",
            "",
            "| Carga | Paneles |",
            "|---:|---:|",
        ]
    )
    lines.extend(f"| {int(load)} | {int(count)} |" for load, count in load_counts.items())
    lines.extend(["", "### Intentos", "", "| Intento | Inscripciones |", "|---:|---:|"])
    lines.extend(f"| {int(attempt)} | {int(count)} |" for attempt, count in retry_counts.items())
    lines.extend(
        [
            "",
            "## Ajustes de factibilidad Beta",
            "",
            f"- Total: {beta_total} de {beta_denominator} contextos ({beta_rate:.3%}).",
            "- El detalle interno se conserva exclusivamente en `Auditoria.csv`.",
            "",
            "## Contraste descriptivo con referencias macro",
            "",
            "Las diferencias siguientes comparan exclusivamente referencias ficticias "
            "con resultados ficticios del mismo demo; no constituyen validación externa.",
            "",
            "| Período | Activos ref. | Activos sim. | Dif. activos | Promedio ref. | Promedio sim. | Dif. promedio |",
            "|---|---:|---:|---:|---:|---:|---:|",
        ]
    )
    refs = bundle.macro_references.loc[
        bundle.macro_references["usar_como_referencia"].eq(True)
    ].copy()
    reference_cohort = (
        int(refs["cohorte"].astype(int).iloc[0]) if len(refs) else None
    )
    if reference_cohort is not None:
        refs = refs.loc[refs["cohorte"].astype(int).eq(reference_cohort)]
        reference_panel = panel.loc[panel["cohorte"].eq(reference_cohort)]
    else:
        reference_panel = panel.iloc[0:0]
    for ref in refs.itertuples(index=False):
        simulated_active = int(reference_panel["periodo"].eq(ref.periodo).sum())
        simulated_semester_means = reference_panel.loc[
            reference_panel["periodo"].eq(ref.periodo), "promedio_semestre"
        ].dropna()
        simulated_mean = (
            float(simulated_semester_means.mean())
            if len(simulated_semester_means)
            else float("nan")
        )
        reference_mean = float(ref.promedio_cohorte_observado)
        lines.append(
            f"| {ref.periodo} | {int(ref.activos_observados)} | {simulated_active} | "
            f"{simulated_active-int(ref.activos_observados):+d} | {reference_mean:.2f} | "
            f"{simulated_mean:.2f} | {simulated_mean-reference_mean:+.2f} |"
        )
    lines.extend(
        [
            "",
            "## Contrato de datos",
            "",
            "- `FEATURE_PRE_T`: información disponible antes de generar el período t.",
            "- `OUTCOME_POST_T` (TARGET): resultados observables sólo después del cierre.",
            "- `AUDIT_ONLY`: estado interno excluido de las cinco tablas públicas.",
            "- La clave multi-seed de estudiante es `(scenario_id, seed, student_id)`.",
            "",
            "## Reproducibilidad",
            "",
            f"- Python: `{result.metadata['python_version']}`",
            f"- RNG: `{result.metadata['rng']['algorithm']}` con derivación `{result.metadata['rng']['derivation']}`.",
            "- `metadata.json` conserva hashes, parámetros, currículo, cohortes, matrices efectivas y contrato.",
            "- `table_hashes.json` permite comprobar los bytes exactos de las seis tablas.",
            "",
        ]
    )
    return "\n".join(lines)


def write_run(
    result: SimulationResult,
    bundle: InputBundle,
    output_root: str | Path,
    central_reports_dir: str | Path | None = None,
) -> Path:
    """Write one seed to an isolated directory; never overwrite an existing run."""

    if result.metadata.get("validation_status", {}).get("structural_validation") != "PASS":
        raise RunVerificationError("La corrida no tiene validación estructural PASS")
    if set(result.tables()) != LOGICAL_TABLE_NAMES:
        raise RunVerificationError("La corrida no contiene exactamente las seis tablas lógicas")

    run_dir = run_directory(
        output_root,
        str(result.metadata["scenario_id"]),
        int(result.metadata["seed"]),
    )
    if run_dir.exists():
        raise FileExistsError(
            f"La corrida ya existe y no se sobrescribe automáticamente: {run_dir}"
        )

    table_payloads = {name: _dataframe_bytes(frame) for name, frame in result.tables().items()}
    table_hashes = {name: _sha256_bytes(payload) for name, payload in table_payloads.items()}
    result.metadata["validation_status"]["strict_json"] = "PASS"
    _apply_release_reporting_status(result.metadata)
    result.metadata["derived_artifacts"] = {
        "analysis_ready_workbook": ANALYSIS_READY_WORKBOOK
    }
    metadata_text = strict_json_dumps(result.metadata) + "\n"
    hashes_text = strict_json_dumps(table_hashes) + "\n"
    summary = build_run_summary(result, bundle, run_dir)

    run_dir.parent.mkdir(parents=True, exist_ok=True)
    run_dir.mkdir()
    for name, payload in table_payloads.items():
        (run_dir / f"{name}.csv").write_bytes(payload)
    (run_dir / "metadata.json").write_text(metadata_text, encoding="utf-8", newline="\n")
    (run_dir / "table_hashes.json").write_text(hashes_text, encoding="utf-8", newline="\n")
    (run_dir / "run_summary.md").write_text(summary, encoding="utf-8", newline="\n")
    write_analysis_ready_workbook(run_dir)

    verify_run_directory(
        run_dir,
        expected_scenario=str(result.metadata["scenario_id"]),
        expected_seed=int(result.metadata["seed"]),
    )

    if central_reports_dir is not None:
        reports_dir = Path(central_reports_dir).resolve()
        reports_dir.mkdir(parents=True, exist_ok=True)
        if int(result.metadata["seed"]) == int(bundle.general_parameters["seed_desarrollo"]):
            scenario_id = str(result.metadata["scenario_id"])
            report_name = f"demo_run_summary_{_safe_scenario(scenario_id)}.md"
            (reports_dir / report_name).write_text(
                summary, encoding="utf-8", newline="\n"
            )
    return run_dir


def verify_run_directory(
    run_dir: str | Path,
    *,
    expected_scenario: str | None = None,
    expected_seed: int | None = None,
) -> dict[str, Any]:
    """Verify completeness, strict JSON, identity and exact CSV hashes for a run."""

    directory = Path(run_dir).resolve()
    if not directory.is_dir():
        raise RunVerificationError(f"No existe el directorio de corrida: {directory}")

    expected_files = {f"{name}.csv" for name in LOGICAL_TABLE_NAMES} | {
        ANALYSIS_READY_WORKBOOK,
        "metadata.json",
        "table_hashes.json",
        "run_summary.md",
    }
    actual_files = {path.name for path in directory.iterdir() if path.is_file()}
    missing = sorted(expected_files - actual_files)
    extra = sorted(actual_files - expected_files)
    if missing or extra:
        raise RunVerificationError(
            f"Contenido de corrida inválido; faltan={missing}, sobran={extra}"
        )

    metadata = strict_json_loads((directory / "metadata.json").read_text(encoding="utf-8"))
    hashes = strict_json_loads((directory / "table_hashes.json").read_text(encoding="utf-8"))
    if not isinstance(metadata, dict) or not isinstance(hashes, dict):
        raise RunVerificationError("metadata.json y table_hashes.json deben ser objetos JSON")
    if set(hashes) != LOGICAL_TABLE_NAMES:
        raise RunVerificationError("table_hashes.json no declara exactamente las seis tablas")
    derived_artifacts = metadata.get("derived_artifacts", {})
    if derived_artifacts != {"analysis_ready_workbook": ANALYSIS_READY_WORKBOOK}:
        raise RunVerificationError("metadata.derived_artifacts no declara el XLSX esperado")

    scenario_id = str(metadata.get("scenario_id", ""))
    try:
        seed = int(metadata.get("seed"))
    except (TypeError, ValueError) as exc:
        raise RunVerificationError("metadata.seed no es un entero válido") from exc
    run_id = f"{scenario_id}_seed_{seed}"
    if metadata.get("run_id") != run_id:
        raise RunVerificationError("metadata.run_id no coincide con scenario_id y seed")
    if expected_scenario is not None and scenario_id != str(expected_scenario):
        raise RunVerificationError(
            f"Escenario inesperado: esperado={expected_scenario}, observado={scenario_id}"
        )
    if expected_seed is not None and seed != int(expected_seed):
        raise RunVerificationError(
            f"Seed inesperada: esperada={int(expected_seed)}, observada={seed}"
        )

    for name in sorted(LOGICAL_TABLE_NAMES):
        table_path = directory / f"{name}.csv"
        actual_hash = _sha256_file(table_path)
        if hashes.get(name) != actual_hash:
            raise RunVerificationError(
                f"Hash inconsistente para {name}: registrado={hashes.get(name)}, actual={actual_hash}"
            )
        frame = pd.read_csv(table_path, low_memory=False)
        if not RUN_KEYS <= set(frame.columns):
            raise RunVerificationError(f"{name}: faltan claves de identidad de corrida")
        if frame.empty or not (
            frame["scenario_id"].astype(str).eq(scenario_id).all()
            and frame["seed"].astype(int).eq(seed).all()
            and frame["run_id"].astype(str).eq(run_id).all()
        ):
            raise RunVerificationError(f"{name}: identidad de corrida inconsistente")
    status = metadata.get("validation_status", {})
    if (
        status.get("structural_validation") != "PASS"
        or status.get("strict_json") != "PASS"
        or status.get("calibration") != "NOT_APPLICABLE_DEMO"
        or status.get("plausibility") != "DEMONSTRATION_ONLY"
    ):
        raise RunVerificationError(
            "La corrida no declara los estados técnicos y finales de reporting esperados"
        )
    if metadata.get("release_status") != RELEASE_STATUS:
        raise RunVerificationError("metadata.release_status no coincide con el demo público")
    if not (directory / "run_summary.md").read_text(encoding="utf-8").strip():
        raise RunVerificationError("run_summary.md está vacío")
    verify_analysis_ready_workbook(directory)

    return {
        "run_dir": str(directory),
        "scenario_id": scenario_id,
        "seed": seed,
        "run_id": run_id,
        "table_hashes": hashes,
        "metadata": metadata,
    }
